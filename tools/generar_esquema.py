from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from datetime import datetime, date
import os

F = "Arial"
HDR = PatternFill("solid", fgColor="1F3864")
KEY = PatternFill("solid", fgColor="BF8F00")
REF = PatternFill("solid", fgColor="2E75B6")
thin = Side(style="thin", color="BFBFBF")

now = datetime(2026,10,15,12,30)
T = {}  # table -> list of (col,type,desc,example)

T["Usuarios"] = [
 ("ID_Usuario","Text (KEY)","Identificador único. AppSheet: UNIQUEID()","U0001"),
 ("Email","Email","Correo con el que inicia sesión (debe ser el mismo de Google/AppSheet)","master@minegocio.com"),
 ("Nombre","Text","Nombre completo","Admin Principal"),
 ("Telefono","Phone","Teléfono","+56911111111"),
 ("ID_Rol","Ref → Roles","Rol base del usuario","R01"),
 ("ID_Delivery","Ref → Deliverys","Solo si el rol es Delivery: vincula al repartidor","" ),
 ("Activo","Yes/No","Si está en FALSE no puede usar la app","TRUE"),
 ("Password_Hash","Text","Hash de la contraseña (lo escribe el backend, nunca se edita a mano)",""),
 ("Salt","Text","Sal aleatoria de la contraseña (la escribe el backend)",""),
 ("Ultimo_Login","DateTime","Último inicio de sesión (lo escribe el backend)",None),
 ("Fecha_Creacion","DateTime","Automático: NOW()",now),
]
T["Roles"] = [
 ("ID_Rol","Text (KEY)","Identificador del rol","R01"),
 ("Nombre_Rol","Text","Master, Ventas, Delivery u otros que crees","Master"),
 ("Descripcion","Text","Para qué sirve el rol","Acceso total a todo el sistema"),
]
T["Permisos"] = [
 ("ID_Permiso","Text (KEY)","Identificador. Fórmula sugerida: [ID_Rol] & \"-\" & [Modulo]","R01-Inventario"),
 ("ID_Rol","Ref → Roles","Rol al que aplica","R01"),
 ("Modulo","Enum: Inventario|Categorias|Ventas|Clientes|Proveedores|Delivery|Historial|Conocimiento|Reportes|KPI|Usuarios|Configuracion","Módulo de la app","Inventario"),
 ("Ver","Yes/No","Puede ver el módulo","TRUE"),
 ("Crear","Yes/No","Puede agregar registros","TRUE"),
 ("Editar","Yes/No","Puede modificar registros","TRUE"),
 ("Eliminar","Yes/No","Puede borrar registros","TRUE"),
 ("Exportar","Yes/No","Puede descargar reportes Excel/PDF","TRUE"),
]
T["Permisos_Usuario"] = [
 ("ID_PermU","Text (KEY)","Identificador","PU0001"),
 ("ID_Usuario","Ref → Usuarios","Usuario con excepción","U0002"),
 ("Modulo","Enum: Inventario|Categorias|Ventas|Clientes|Proveedores|Delivery|Historial|Conocimiento|Reportes|KPI|Usuarios|Configuracion","Módulo","Ventas"),
 ("Ver","Yes/No","Sobrescribe el permiso del rol","TRUE"),
 ("Crear","Yes/No","Sobrescribe el permiso del rol","TRUE"),
 ("Editar","Yes/No","Sobrescribe el permiso del rol","TRUE"),
 ("Eliminar","Yes/No","Sobrescribe el permiso del rol","FALSE"),
 ("Exportar","Yes/No","Sobrescribe el permiso del rol","FALSE"),
]
T["Configuracion"] = [
 ("Clave","Text (KEY)","Nombre del parámetro","NOMBRE_NEGOCIO"),
 ("Valor","Text","Valor del parámetro","Mi Negocio"),
 ("Descripcion","Text","Qué hace","Nombre que sale en boletas y facturas"),
]
T["Categorias"] = [
 ("ID_Categoria","Text (KEY)","Identificador. AppSheet: UNIQUEID()","C001"),
 ("Nombre","Text (LABEL)","Nombre de la categoría","Bebidas"),
 ("Imagen","Image","Imagen que se muestra en el selector de venta","Categorias_Images/bebidas.png"),
 ("Orden","Number","Orden de aparición en pantalla",1),
 ("Activo","Yes/No","Mostrar u ocultar","TRUE"),
]
T["Unidades"] = [
 ("Unidad","Text (KEY)","Unidad de medida (lista base para el Enum)","Unidad"),
 ("Abreviatura","Text","Abreviatura","un"),
]
T["Productos"] = [
 ("ID_Producto","Text (KEY)","Identificador. AppSheet: UNIQUEID()","P0001"),
 ("Codigo_Barras","Text","Se llena con el escáner de la cámara (tipo Text con Scan)","7801234567890"),
 ("SKU","Text","Código interno opcional","BEB-001"),
 ("Nombre","Text (LABEL)","Nombre del producto","Bebida cola 1.5L"),
 ("ID_Categoria","Ref → Categorias","Categoría","C001"),
 ("Imagen","Image","Foto del producto","Productos_Images/cola15.jpg"),
 ("Unidad_Medida","Enum (lista = Unidades)","Unidad, kg, litro, caja, etc.","Unidad"),
 ("Precio_Compra","Price","Costo de compra","900"),
 ("Precio_Venta","Price","Precio de venta","1500"),
 ("Stock_Actual","Decimal","Se descuenta automáticamente al vender","48"),
 ("Stock_Minimo","Decimal","Al llegar a este número o menos se genera alerta","10"),
 ("ID_Proveedor","Ref → Proveedores","Proveedor habitual","PR001"),
 ("Activo","Yes/No","Producto disponible para venta","TRUE"),
 ("Fecha_Creacion","DateTime","Automático",now),
 ("Va_A_Cocina","Yes/No","Si es TRUE (o vacío) el producto sale en la comanda de cocina","TRUE"),
]
T["Movimientos_Inventario"] = [
 ("ID_Mov","Text (KEY)","Identificador","M00001"),
 ("Fecha","DateTime","Automático",now),
 ("ID_Producto","Ref → Productos","Producto afectado","P0001"),
 ("Tipo","Enum: Entrada compra|Salida venta|Ajuste|Merma|Anulacion venta|Anulacion compra","Motivo del movimiento","Salida venta"),
 ("Cantidad","Decimal","Positiva siempre; el Tipo define si suma o resta","2"),
 ("Stock_Antes","Decimal","Stock antes del movimiento","50"),
 ("Stock_Despues","Decimal","Stock luego del movimiento","48"),
 ("Referencia","Text","ID_Venta o ID_Compra que lo originó","V000001"),
 ("Usuario","Email","Quién lo hizo: USEREMAIL()","master@minegocio.com"),
 ("Nota","Text","Observación","" ),
]
T["Clientes"] = [
 ("ID_Cliente","Text (KEY)","Identificador","CL0001"),
 ("Nombre","Text (LABEL)","Nombre o razón social","María González"),
 ("Documento","Text","RUT / DNI / RIF","12.345.678-9"),
 ("Telefono","Phone","Teléfono","+56922222222"),
 ("Email","Email","Correo (para enviar boletas/facturas)","maria@correo.com"),
 ("Direccion","Address","Dirección principal","Av. Principal 123"),
 ("Ubicacion","LatLong","Coordenadas para el mapa","-33.4489,-70.6693"),
 ("Notas","LongText","Observaciones","Prefiere entrega en la tarde"),
 ("Fecha_Creacion","DateTime","Automático",now),
]
T["Proveedores"] = [
 ("ID_Proveedor","Text (KEY)","Identificador","PR001"),
 ("Nombre","Text (LABEL)","Empresa o persona","Distribuidora Norte"),
 ("Documento","Text","RUT / ID fiscal","76.111.222-3"),
 ("Contacto","Text","Persona de contacto","Carlos Soto"),
 ("Telefono","Phone","Teléfono","+56933333333"),
 ("Email","Email","Correo","ventas@distnorte.com"),
 ("Direccion","Address","Dirección","Calle Industrial 45"),
 ("Dias_Visita","EnumList: Lunes|Martes|Miercoles|Jueves|Viernes|Sabado|Domingo","Días en que visita tu local","Martes, Viernes"),
 ("Ultima_Visita","Date","Última visita registrada","2026-10-13"),
 ("Proxima_Visita","Date","Próxima visita estimada","2026-10-16"),
 ("Notas","LongText","Condiciones, plazos, etc.","Pago a 15 días"),
]
T["Visitas_Proveedor"] = [
 ("ID_Visita","Text (KEY)","Identificador","VP0001"),
 ("ID_Proveedor","Ref → Proveedores","Proveedor","PR001"),
 ("Fecha","Date","Fecha de la visita",date(2026,10,13)),
 ("Tipo","Enum: Proveedor visita el local|Yo voy al proveedor","Quién se desplaza","Proveedor visita el local"),
 ("Notas","LongText","Detalle (qué se pidió, qué se pagó)","Pedido semanal"),
 ("Usuario","Email","Quién lo registró","master@minegocio.com"),
]
T["Compras"] = [
 ("ID_Compra","Text (KEY)","Identificador","CO0001"),
 ("Fecha","DateTime","Automático",now),
 ("ID_Proveedor","Ref → Proveedores","Proveedor","PR001"),
 ("Total","Price","Suma de Detalle_Compra (virtual column en AppSheet)","43200"),
 ("Estado","Enum: Pendiente|Recibida|Anulada","Al pasar a Recibida suma stock","Recibida"),
 ("Usuario","Email","Quién registró","master@minegocio.com"),
 ("Nota","Text","Observación (N° de factura del proveedor, etc.)",""),
]
T["Detalle_Compra"] = [
 ("ID_DetCompra","Text (KEY)","Identificador","DC0001"),
 ("ID_Compra","Ref → Compras (Is a part of)","Compra a la que pertenece","CO0001"),
 ("ID_Producto","Ref → Productos","Producto","P0001"),
 ("Cantidad","Decimal","Cantidad recibida","48"),
 ("Costo_Unitario","Price","Costo por unidad","900"),
 ("Subtotal","Price","Cantidad × Costo (fórmula App formula)","43200"),
]
T["Deliverys"] = [
 ("ID_Delivery","Text (KEY)","Identificador","D001"),
 ("Nombre","Text (LABEL)","Nombre del repartidor","Juan Pérez"),
 ("Telefono","Phone","Teléfono","+56944444444"),
 ("Email","Email","Correo con el que entra a la app (perfil Delivery)","juan.delivery@correo.com"),
 ("Vehiculo","Enum: Moto|Bicicleta|Auto|A pie","Medio de transporte","Moto"),
 ("Patente","Text","Patente o placa","ABCD12"),
 ("Estado","Enum: Disponible|Ocupado|Inactivo","Estado actual (cambia solo con automatización)","Disponible"),
 ("Tarifa_Base","Price","Tarifa base por servicio","1000"),
 ("Tarifa_Km","Price","Tarifa por km (opcional)","300"),
 ("Activo","Yes/No","Habilitado","TRUE"),
 ("Ultima_Ubicacion","LatLong","Última posición reportada","-33.4500,-70.6600"),
 ("Fecha_Ubicacion","DateTime","Hora de la última posición",now),
]
T["Ventas"] = [
 ("ID_Venta","Text (KEY)","Identificador. AppSheet: UNIQUEID() o correlativo","V000001"),
 ("Fecha_Hora","DateTime","Automático: NOW()",now),
 ("ID_Cliente","Ref → Clientes","Cliente (se puede crear desde la venta con Add new)","CL0001"),
 ("Tipo_Entrega","Enum: Retiro en local|Para servir|Delivery","Si es Delivery se crea un servicio","Delivery"),
 ("Subtotal","Price","Suma de Detalle_Venta (virtual column)","3000"),
 ("Descuento","Price","Descuento total","0"),
 ("Costo_Delivery","Price","Monto del delivery (si aplica)","1500"),
 ("Total","Price","Subtotal - Descuento + Costo_Delivery (virtual)","4500"),
 ("Metodo_Pago","Enum: Efectivo|Tarjeta|Transferencia|Mixto","Forma de pago","Efectivo"),
 ("Estado_Pago","Enum: Pagado|Pendiente|Parcial","Estado del pago","Pagado"),
 ("Estado","Enum: Borrador|Confirmada|En despacho|Entregada|Anulada","Estado de la venta","Confirmada"),
 ("Tipo_Documento","Enum: Ticket|Boleta|Factura","Documento a emitir","Boleta"),
 ("Dia_Semana","Text","Virtual/Apps Script: TEXT([Fecha_Hora],\"dddd\"). Sirve para KPI","jueves"),
 ("Usuario","Email","Vendedor: USEREMAIL()","ventas@minegocio.com"),
 ("Referencia","Text","Mesa, nombre o número para identificar el pedido en cocina","Mesa 4"),
 ("Nota_Cocina","Text","Indicación general para la cocina","Todo para llevar junto"),
]
T["Detalle_Venta"] = [
 ("ID_Detalle","Text (KEY)","Identificador","DV000001"),
 ("ID_Venta","Ref → Ventas (Is a part of)","Venta a la que pertenece","V000001"),
 ("ID_Producto","Ref → Productos","Producto vendido","P0001"),
 ("Nombre_Producto","Text","Copia del nombre al momento de vender (historial fiel)","Bebida cola 1.5L"),
 ("ID_Categoria","Ref → Categorias","Copia de categoría para KPI por categoría","C001"),
 ("Cantidad","Decimal","Cantidad vendida","2"),
 ("Precio_Unitario","Price","Precio al momento de la venta","1500"),
 ("Costo_Unitario","Price","Costo al momento de la venta (para margen)","900"),
 ("Descuento","Price","Descuento de la línea","0"),
 ("Subtotal","Price","Cantidad × Precio - Descuento","3000"),
 ("Fecha","Date","Fecha de la venta (para filtros y KPI)",date(2026,10,15)),
 ("Nota","Text","Indicación para cocina de este producto (ej. sin cebolla)",""),
 ("Va_A_Cocina","Yes/No","Copia del producto: sale en la comanda","TRUE"),
]
T["Servicios_Delivery"] = [
 ("ID_Servicio","Text (KEY)","Identificador","S000001"),
 ("ID_Venta","Ref → Ventas","Venta asociada","V000001"),
 ("ID_Delivery","Ref → Deliverys","Repartidor asignado","D001"),
 ("Nombre_Delivery","Text","Copia del nombre (historial fiel aunque cambie después)","Juan Pérez"),
 ("Fecha_Asignacion","DateTime","Automático: NOW() al asignar",now),
 ("Fecha_Aceptado","DateTime","Se llena al presionar Aceptar",None),
 ("Fecha_Finalizado","DateTime","Se llena al presionar Finalizar",None),
 ("Fecha","Date","Solo la fecha, para filtrar por rango",date(2026,10,15)),
 ("Direccion_Entrega","Address","Dirección de destino","Av. Principal 123"),
 ("Destino_LatLong","LatLong","Coordenadas destino","-33.4489,-70.6693"),
 ("Origen_LatLong","LatLong","Coordenadas del local","-33.4372,-70.6506"),
 ("Distancia_Km","Decimal","Km (manual o calculado con Apps Script/Maps)","4.2"),
 ("Tiempo_Min","Number","Duración en minutos","18"),
 ("Monto_Delivery","Price","Costo del servicio","1500"),
 ("Estado","Enum: Asignado|Aceptado|En camino|Finalizado|Rechazado|Cancelado","Estado del servicio","Asignado"),
 ("Notas","LongText","Instrucciones para el repartidor","Llamar al llegar"),
 ("Usuario_Asigna","Email","Quién asignó","ventas@minegocio.com"),
]
T["Ubicaciones_Delivery"] = [
 ("ID_Ubic","Text (KEY)","Identificador","UB000001"),
 ("ID_Servicio","Ref → Servicios_Delivery","Servicio en curso","S000001"),
 ("ID_Delivery","Ref → Deliverys","Repartidor","D001"),
 ("LatLong","LatLong","Posición registrada (AppSheet: HERE())","-33.4470,-70.6620"),
 ("Fecha_Hora","DateTime","Momento del registro",now),
]
T["Documentos"] = [
 ("ID_Documento","Text (KEY)","Identificador","DOC0001"),
 ("ID_Venta","Ref → Ventas","Venta origen","V000001"),
 ("Tipo","Enum: Ticket|Boleta|Factura","Tipo de documento","Boleta"),
 ("Folio","Text","Número correlativo interno (o folio SII si usas facturación electrónica)","1001"),
 ("Fecha","DateTime","Emisión",now),
 ("Neto","Price","Monto neto","3782"),
 ("Impuesto","Price","IVA u otro impuesto","718"),
 ("Total","Price","Total","4500"),
 ("Formato","Enum: Termica 58mm|Termica 80mm|A4 Carta","Formato de impresión","Termica 80mm"),
 ("PDF_URL","File","PDF generado por Apps Script en Drive",""),
]
T["Categorias_Conocimiento"] = [
 ("ID_CatCon","Text (KEY)","Identificador","CK01"),
 ("Nombre","Text (LABEL)","Nombre de la categoría","Recetas"),
 ("Imagen","Image","Imagen opcional","Conocimiento_Images/recetas.png"),
]
T["Conocimiento"] = [
 ("ID_Conocimiento","Text (KEY)","Identificador","K0001"),
 ("Titulo","Text (LABEL)","Título del documento","Preparación salsa base"),
 ("ID_CatCon","Ref → Categorias_Conocimiento","Categoría","CK01"),
 ("Archivo_PDF","File","PDF subido (se guarda en Drive)","Conocimiento_Files/salsa_base.pdf"),
 ("Descripcion","LongText","Resumen","Receta estándar de salsa base"),
 ("Etiquetas","EnumList (Allow other values)","Etiquetas para buscar","salsa, base"),
 ("Fecha","DateTime","Automático",now),
 ("Usuario","Email","Quién lo subió","master@minegocio.com"),
]
T["Alertas"] = [
 ("ID_Alerta","Text (KEY)","Identificador","A000001"),
 ("Fecha","DateTime","Automático",now),
 ("Tipo","Enum: Stock minimo|Delivery finalizado|Delivery rechazado|Otro","Tipo de alerta","Stock minimo"),
 ("Referencia","Text","ID_Producto o ID_Servicio","P0001"),
 ("Mensaje","Text","Texto de la alerta","Bebida cola 1.5L llegó al stock mínimo (10)"),
 ("Para_Rol","Ref → Roles","Rol que debe verla","R01"),
 ("Leida","Yes/No","Marcada como leída","FALSE"),
]
T["Auditoria"] = [
 ("ID_Log","Text (KEY)","Identificador","L000001"),
 ("Fecha","DateTime","Automático",now),
 ("Usuario","Email","Quién","master@minegocio.com"),
 ("Modulo","Text","Módulo afectado","Inventario"),
 ("Accion","Enum: Crear|Editar|Eliminar|Exportar|Login","Qué hizo","Crear"),
 ("Registro","Text","ID del registro afectado","P0001"),
 ("Detalle","LongText","Detalle del cambio","Producto creado"),
]

# seed rows
modules = ["Inventario","Categorias","Ventas","Clientes","Proveedores","Delivery","Historial","Conocimiento","Reportes","KPI","Usuarios","Configuracion"]
roles = [("R01","Master","Acceso total"),("R02","Ventas","Vende, consulta inventario, asigna delivery"),("R03","Delivery","Solo ve y gestiona sus servicios")]
perm = {  # role -> module -> (ver,crear,editar,eliminar,exportar)
 "R01": {m:(True,True,True,True,True) for m in modules},
 "R02": {"Inventario":(True,True,True,False,False),"Categorias":(True,False,False,False,False),"Ventas":(True,True,True,False,False),
         "Clientes":(True,True,True,False,False),"Proveedores":(True,False,False,False,False),"Delivery":(True,True,False,False,False),
         "Historial":(True,False,False,False,False),"Conocimiento":(True,False,False,False,False),"Reportes":(False,False,False,False,False),
         "KPI":(False,False,False,False,False),"Usuarios":(False,False,False,False,False),"Configuracion":(False,False,False,False,False)},
 "R03": {m:(False,False,False,False,False) for m in modules},
}
perm["R03"]["Delivery"]=(True,False,True,False,False)
perm["R03"]["Historial"]=(True,False,False,False,False)
perm["R03"]["Conocimiento"]=(True,False,False,False,False)

seed = {}
seed["Roles"] = [list(r) for r in roles]
seed["Permisos"] = [[f"{r}-{m}",r,m,*perm[r][m]] for r in perm for m in modules]
seed["Unidades"] = [[u,a] for u,a in [("Unidad","un"),("Kilogramo","kg"),("Gramo","g"),("Litro","lt"),("Mililitro","ml"),("Caja","caja"),("Paquete","paq"),("Docena","doc")]]
seed["Configuracion"] = [
 ["NOMBRE_NEGOCIO","Mi Negocio","Nombre que sale en documentos"],
 ["DOCUMENTO_NEGOCIO","76.000.000-0","RUT / ID fiscal del negocio"],
 ["DIRECCION_NEGOCIO","Calle Ejemplo 100","Dirección del local"],
 ["MONEDA","CLP","Moneda"],
 ["IMPUESTO_PORC","0.19","Porcentaje de impuesto (0.19 = 19%)"],
 ["TARIFA_BASE_DELIVERY","1000","Tarifa base por defecto"],
 ["TARIFA_KM_DELIVERY","300","Tarifa por km por defecto"],
 ["FORMATO_DOC_DEFECTO","Termica 80mm","Termica 58mm | Termica 80mm | A4 Carta"],
 ["FORMATO_COMANDA","Termica 80mm","Formato del ticket de cocina"],
 ["COMANDA_AUTO","FALSE","TRUE = al vender se abre sola la impresión de la comanda"],
 ["ORIGEN_LATLONG","-33.4372,-70.6506","Coordenadas del local"],
 ["ALERTA_STOCK_ACTIVA","TRUE","Activa las alertas de stock mínimo"],
 ["LOGO","","URL del logo (se sube desde Configuración)"],
]

wb = Workbook()
ws0 = wb.active; ws0.title = "LEEME"

def style_header(ws, cols):
    for i,(name,typ,desc,ex) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = KEY if i==1 else (REF if typ.startswith("Ref") else HDR)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = max(14, min(32, len(name)+6))
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "B2"

for name, cols in T.items():
    ws = wb.create_sheet(name)
    style_header(ws, cols)
    rows = seed.get(name) or [[c[3] for c in cols]]
    for r, row in enumerate(rows, start=2):
        for ci, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=ci, value=v)
            cell.font = Font(name=F)
            if isinstance(v, datetime): cell.number_format = "dd-mm-yyyy hh:mm"
            elif isinstance(v, date): cell.number_format = "dd-mm-yyyy"
    # list validations for enums
    last = max(len(rows)+1, 500)
    for i,(cn,typ,desc,ex) in enumerate(cols, start=1):
        if typ.startswith("Enum:") or typ.startswith("EnumList:"):
            opts = typ.split(":",1)[1].strip()
            if typ.startswith("EnumList:"): continue
            if len(opts) < 250:
                dv = DataValidation(type="list", formula1=f'"{",".join(o.strip() for o in opts.split("|"))}"', allow_blank=True)
                ws.add_data_validation(dv); dv.add(f"{get_column_letter(i)}2:{get_column_letter(i)}{last}")
        if typ == "Yes/No":
            dv = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
            ws.add_data_validation(dv); dv.add(f"{get_column_letter(i)}2:{get_column_letter(i)}{last}")

# DICCIONARIO
wd = wb.create_sheet("DICCIONARIO", 1)
hd = ["Tabla","Columna","Tipo en AppSheet","Descripción / regla","Ejemplo"]
for i,h in enumerate(hd, start=1):
    c = wd.cell(row=1,column=i,value=h); c.font=Font(name=F,bold=True,color="FFFFFF"); c.fill=HDR
    c.alignment=Alignment(horizontal="center")
r=2
for tname, cols in T.items():
    for cn,typ,desc,ex in cols:
        vals=[tname,cn,typ,desc,"" if ex is None else (ex.strftime("%d-%m-%Y %H:%M") if isinstance(ex,datetime) else ex.strftime("%d-%m-%Y") if isinstance(ex,date) else str(ex))]
        for ci,v in enumerate(vals,start=1):
            c=wd.cell(row=r,column=ci,value=v); c.font=Font(name=F, bold=(ci==1 and cn==cols[0][0]))
            c.alignment=Alignment(wrap_text=True,vertical="top")
        r+=1
for col,w in zip("ABCDE",[24,22,38,70,32]): wd.column_dimensions[col].width=w
wd.freeze_panes="A2"; wd.auto_filter.ref=f"A1:E{r-1}"

# LEEME
lines = [
 ("BASE DE DATOS – INVENTARIO, VENTAS, DELIVERY", "t"),
 ("", ""),
 ("CÓMO USAR ESTE ARCHIVO", "h"),
 ("1. Súbelo a Google Drive y ábrelo con Google Sheets (Archivo > Guardar como Hoja de cálculo de Google).", ""),
 ("2. Cada pestaña es una tabla. La FILA 1 son los encabezados y NO se deben renombrar ni mover (AppSheet los usa como nombres de columna).", ""),
 ("3. La FILA 2 de cada pestaña es un registro de EJEMPLO para que AppSheet detecte los tipos. Bórrala (o reemplázala) antes de usar en producción.", ""),
 ("4. Esta plantilla es de REFERENCIA: la app crea todas estas hojas sola al ejecutar instalar() en Apps Script (ver README del repositorio). No necesitas subir este archivo.", ""),
 ("5. La pestaña DICCIONARIO explica cada columna, su tipo en AppSheet y las relaciones (Ref).", ""),
 ("", ""),
 ("COLORES DE ENCABEZADO", "h"),
 ("Dorado = columna llave (KEY, valor único).   Azul claro = columna de relación (Ref a otra tabla).   Azul oscuro = columna normal.", ""),
 ("", ""),
 ("MÓDULOS Y TABLAS", "h"),
 ("Inventario: Productos, Categorias, Unidades, Movimientos_Inventario, Compras, Detalle_Compra", ""),
 ("Ventas: Ventas, Detalle_Venta, Documentos", ""),
 ("Clientes: Clientes", ""),
 ("Proveedores: Proveedores, Visitas_Proveedor", ""),
 ("Delivery: Deliverys, Servicios_Delivery, Ubicaciones_Delivery  (el HISTORIAL es una vista filtrada de Servicios_Delivery, no una tabla aparte)", ""),
 ("Usuarios y permisos: Usuarios, Roles, Permisos, Permisos_Usuario", ""),
 ("Base de conocimiento: Categorias_Conocimiento, Conocimiento", ""),
 ("Sistema: Configuracion, Alertas, Auditoria", ""),
 ("", ""),
 ("REGLAS IMPORTANTES", "h"),
 ("• Los valores Sí/No se guardan como TRUE/FALSE.", ""),
 ("• Las columnas ID las genera AppSheet con UNIQUEID() (Initial value). No las escribas a mano en la app.", ""),
 ("• No uses fórmulas de Sheets dentro de estas tablas: los totales y cálculos se hacen en AppSheet (App formula / Virtual columns).", ""),
 ("• Imágenes y PDFs se guardan en carpetas de Drive que AppSheet crea junto al Sheet; en la celda queda solo la ruta.", ""),
]
for i,(t,k) in enumerate(lines, start=1):
    c = ws0.cell(row=i, column=1, value=t)
    c.font = Font(name=F, bold=(k in ("t","h")), size=14 if k=="t" else 11, color="1F3864" if k in("t","h") else "000000")
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws0.column_dimensions["A"].width = 130

wb.save(os.path.join(os.path.dirname(os.path.abspath(__file__)),"..","docs","BaseDatos_Plantilla.xlsx"))
print(len(T), "tablas;", sum(len(v) for v in T.values()), "columnas")

# ---- export schema + seeds for Apps Script ----
import json, os
os.makedirs(os.path.join(os.path.dirname(os.path.abspath(__file__)),"..","backend"), exist_ok=True)
schema = {k:[c[0] for c in v] for k,v in T.items()}
seeds = {k:[[ (x if isinstance(x,(str,int,float,bool)) or x is None else str(x)) for x in row] for row in rows] for k,rows in seed.items()}
with open(os.path.join(os.path.dirname(os.path.abspath(__file__)),"..","backend","Schema.gs"),"w") as f:
    f.write("// GENERADO desde la base de datos. Define las hojas (tablas) y sus columnas.\n")
    f.write("const SCHEMA = "+json.dumps(schema,ensure_ascii=False,indent=1)+";\n\n")
    f.write("const SEED = "+json.dumps(seeds,ensure_ascii=False,indent=1)+";\n")
print("Schema.gs written")

# ---- archivo único para pegar en Apps Script ----
base = os.path.join(os.path.dirname(os.path.abspath(__file__)),"..","backend")
parts = [open(os.path.join(base,f),encoding="utf8").read() for f in ("Config.gs","Schema.gs","Code.gs","Setup.gs")]
head = "// ============================================================\n//  CÓDIGO COMPLETO DEL SERVIDOR (un solo archivo para pegar en Apps Script)\n//  Contiene: Config + Schema + Code + Setup. Edita ADMIN_EMAIL y ADMIN_PASSWORD abajo.\n// ============================================================\n\n"
open(os.path.join(base,"TODO_EN_UN_ARCHIVO.gs"),"w",encoding="utf8").write(head+"\n\n".join(parts))
print("TODO_EN_UN_ARCHIVO.gs written")
